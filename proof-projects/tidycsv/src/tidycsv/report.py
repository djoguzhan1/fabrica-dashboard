"""Write the cleaned data and a summary to an Excel workbook."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .clean import CleanResult

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
MUTED_FONT = Font(color="666666")
NUMBER_FORMAT = "#,##0.00"
INT_FORMAT = "#,##0"
DATE_FORMAT = "yyyy-mm-dd"


def _style_header(ws: Worksheet, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _print_setup(ws: Worksheet) -> None:
    """Landscape, fit to one page wide: clients do print these."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _autosize(ws: Worksheet, min_width: int = 8, max_width: int = 48) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value)) if not isinstance(cell.value, datetime) else 10
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, width + 2))


def _is_integer_series(series: pd.Series) -> bool:
    values = series.dropna()
    return not values.empty and bool((values == values.round()).all())


def write_data_sheet(ws: Worksheet, frame: pd.DataFrame, column_types: dict[str, str]) -> None:
    ws.title = "Data"
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        ws.append([None if pd.isna(v) else (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in row])

    ncols = len(frame.columns)
    nrows = len(frame) + 1
    _style_header(ws, 1, ncols)
    for idx, col in enumerate(frame.columns, start=1):
        kind = column_types.get(col)
        fmt = None
        if kind == "date":
            fmt = DATE_FORMAT
        elif kind == "number":
            fmt = INT_FORMAT if _is_integer_series(frame[col]) else NUMBER_FORMAT
        if fmt:
            for r in range(2, nrows + 1):
                ws.cell(row=r, column=idx).number_format = fmt
    if nrows > 1:
        table = Table(displayName="CleanData", ref=f"A1:{get_column_letter(ncols)}{nrows}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    _autosize(ws)
    _print_setup(ws)


def _column_stats(frame: pd.DataFrame, column_types: dict[str, str]) -> list[list[object]]:
    rows: list[list[object]] = []
    for col in frame.columns:
        series = frame[col]
        kind = column_types.get(col, "text")
        non_null = int(series.notna().sum())
        row: list[object] = [col, kind, non_null, int(series.isna().sum()), int(series.nunique(dropna=True))]
        top = series.dropna().astype(str).value_counts()
        row.append(top.index[0] if not top.empty and kind == "text" else None)
        if kind == "number" and non_null:
            row += [float(series.min()), float(series.max()), float(series.sum()), float(series.mean())]
        elif kind == "date" and non_null:
            row += [series.min().to_pydatetime(), series.max().to_pydatetime(), None, None]
        else:
            row += [None, None, None, None]
        rows.append(row)
    return rows


def pick_group_columns(frame: pd.DataFrame, column_types: dict[str, str]) -> tuple[str | None, str | None]:
    """Choose a sensible category column and a numeric column for the chart."""
    category = None
    for col, kind in column_types.items():
        if kind == "text" and 2 <= frame[col].nunique(dropna=True) <= 12:
            category = col
            break
    numeric = next((c for c, k in column_types.items() if k == "number" and not c.endswith("_id") and c != "id"), None)
    return category, numeric


def write_summary_sheet(
    ws: Worksheet,
    result: CleanResult,
    source_name: str,
    group_by: str | None,
    sum_column: str | None,
) -> None:
    ws.title = "Summary"
    frame = result.frame

    ws["A1"] = "Data report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Source: {source_name}"
    ws["A2"].font = MUTED_FONT
    ws["A3"] = f"Generated: {datetime.now():%Y-%m-%d %H:%M}"
    ws["A3"].font = MUTED_FONT

    ws["A5"] = "Rows in"
    ws["B5"] = result.rows_in
    ws["A6"] = "Rows out"
    ws["B6"] = result.rows_out
    ws["A7"] = "Rows removed"
    ws["B7"] = result.rows_in - result.rows_out
    ws["A8"] = "Columns"
    ws["B8"] = len(frame.columns)
    for r in range(5, 9):
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.cell(row=r, column=2).number_format = INT_FORMAT

    header_row = 10
    headers = ["Column", "Type", "Non-empty", "Empty", "Unique", "Most common", "Min", "Max", "Sum", "Mean"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    _style_header(ws, header_row, len(headers))
    for i, row in enumerate(_column_stats(frame, result.column_types), start=header_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=c, value=value)
            if isinstance(value, float):
                cell.number_format = NUMBER_FORMAT
            elif isinstance(value, datetime):
                cell.number_format = DATE_FORMAT
            elif isinstance(value, int) and c in (3, 4, 5):
                cell.number_format = INT_FORMAT
            if c == 6:
                cell.alignment = Alignment(indent=1)

    next_row = header_row + len(frame.columns) + 3

    if group_by is None or sum_column is None:
        auto_cat, auto_num = pick_group_columns(frame, result.column_types)
        group_by = group_by or auto_cat
        sum_column = sum_column or auto_num

    if group_by and sum_column and group_by in frame.columns and sum_column in frame.columns:
        grouped = (
            frame.groupby(group_by, dropna=True)[sum_column]
            .agg(["count", "sum"])
            .sort_values("sum", ascending=False)
            .reset_index()
        )
        ws.cell(row=next_row, column=1, value=f"{sum_column} by {group_by}").font = Font(bold=True, size=12)
        table_header = next_row + 1
        for c, h in enumerate([group_by, "Rows", f"Total {sum_column}"], start=1):
            ws.cell(row=table_header, column=c, value=h)
        _style_header(ws, table_header, 3)
        for i, row in enumerate(grouped.itertuples(index=False), start=table_header + 1):
            ws.cell(row=i, column=1, value=str(row[0]))
            ws.cell(row=i, column=2, value=int(row[1])).number_format = INT_FORMAT
            ws.cell(row=i, column=3, value=float(row[2])).number_format = NUMBER_FORMAT
        last = table_header + len(grouped)

        chart = BarChart()
        chart.type = "col"
        chart.title = f"Total {sum_column} by {group_by}"
        chart.y_axis.title = sum_column
        chart.x_axis.title = group_by
        chart.legend = None
        chart.x_axis.tickLblPos = "low"  # keep labels under the axis when there are negative bars
        chart.height = 8
        chart.width = 18
        data = Reference(ws, min_col=3, min_row=table_header, max_row=last)
        categories = Reference(ws, min_col=1, min_row=table_header + 1, max_row=last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        ws.add_chart(chart, f"E{next_row}")

    _autosize(ws)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 22)
    _print_setup(ws)


def write_changes_sheet(ws: Worksheet, result: CleanResult) -> None:
    ws.title = "Changes"
    headers = ["Where", "What", "Count", "Detail"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    if not result.changes:
        ws.append(["(table)", "Nothing needed changing", 0, ""])
    for change in result.changes:
        ws.append(list(change.as_row()))
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=3).number_format = INT_FORMAT
    ws.freeze_panes = "A2"
    _autosize(ws, max_width=70)
    _print_setup(ws)


def write_report(
    result: CleanResult,
    output: Path,
    *,
    source_name: str,
    group_by: str | None = None,
    sum_column: str | None = None,
) -> Path:
    wb = Workbook()
    write_data_sheet(wb.active, result.frame, result.column_types)
    write_summary_sheet(wb.create_sheet(), result, source_name, group_by, sum_column)
    write_changes_sheet(wb.create_sheet(), result)
    # Open on the summary; that is what the client wants to see first.
    wb.active = 1
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
