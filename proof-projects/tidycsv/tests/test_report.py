from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from tidycsv.clean import clean
from tidycsv.cli import main
from tidycsv.report import pick_group_columns, write_report

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "orders_messy.csv"


def test_report_has_three_sheets(tmp_path):
    frame = pd.DataFrame(
        {
            "Region": ["North", "South", "North", "West"],
            "Amount": ["$10", "$20", "$30", "$40"],
            "Date": ["2026-01-01", "2026-01-02", "2026-01-03", "2026-01-04"],
        }
    )
    result = clean(frame)
    out = write_report(result, tmp_path / "out.xlsx", source_name="test.csv")

    wb = load_workbook(out)
    assert wb.sheetnames == ["Data", "Summary", "Changes"]

    data = wb["Data"]
    assert [c.value for c in data[1]] == ["region", "amount", "date"]
    assert data.max_row == 5
    assert data.freeze_panes == "A2"
    assert data.cell(row=2, column=3).number_format == "yyyy-mm-dd"

    summary = wb["Summary"]
    assert summary["B5"].value == 4
    assert summary["B6"].value == 4
    assert len(summary._charts) == 1  # noqa: SLF001 - openpyxl keeps charts here

    changes = wb["Changes"]
    assert changes.max_row >= 2


def test_pick_group_columns_skips_ids():
    frame = pd.DataFrame({"order_id": ["1", "2"], "region": ["a", "b"], "amount": [1.0, 2.0]})
    types = {"order_id": "number", "region": "text", "amount": "number"}
    assert pick_group_columns(frame, types) == ("region", "amount")


def test_cli_on_example_file(tmp_path, capsys):
    out = tmp_path / "report.xlsx"
    code = main([str(EXAMPLE), "-o", str(out), "--group-by", "Region", "--sum", "Amount ($)"])
    assert code == 0
    assert out.exists()
    printed = capsys.readouterr().out
    assert "rows in" in printed and "wrote" in printed

    wb = load_workbook(out)
    summary = wb["Summary"]
    titles = [cell.value for row in summary.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert "amount by region" in titles


def test_cli_rejects_unknown_column(tmp_path, capsys):
    code = main([str(EXAMPLE), "-o", str(tmp_path / "x.xlsx"), "--group-by", "nope"])
    assert code == 2
    assert "not found" in capsys.readouterr().err
